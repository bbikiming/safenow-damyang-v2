#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
분류기준 생성기 — 담양군_2025_문서분류_대시보드.xlsx 「기준표」 → js/doc-taxonomy-data.js

  (구 이름 build-taxonomy-v33.py — v4.0 부터 버전을 파일명에 박지 않는다. 버전은
   엑셀 「대시보드」 시트 A2 «기준 vX.Y (…)» 에서 읽어 META.version 에 싣는다.)

  v3.3 → v4.0 (2026-09-07) — 적용 법령을 2종으로 좁혔다
  ──────────────────────────────────────────────────
  2026-08-21 군청 미팅에서 이 사업의 적용 법령을 **중대재해처벌법 · 산업안전보건법**
  으로 한정했다. 그 결과 재난안전법(DSM 48항목)·개별 시설법(FAC 16항목)·관리업무
  MGT-03~05 가 통째로 빠져 **이행항목 95 → 28 · 업무단계 213 → 88** 이 됐다.
  살아남은 88단계는 코드가 전부 그대로다(바뀐 것은 OSH-09-03 단계명 · OSH-03-03
  법령근거 2건뿐). 그래서 문서 매핑(stageIds 는 코드다)과 우리 판단(수행경로·
  완료판정 …)이 코드로 무수정 이어진다.

  빠진 125단계에 붙어 있던 문서는 원장 쪽에서 「적용 법령 밖」 제외로 바뀐다 —
  이 파일이 아니라 build-ledger-2025.py 가 그 사실을 싣는다. **삭제 단계를 이 파일에
  숨겨 두지 않는다**(2026-09-08 확정 — 완전 교체). 컨설팅 회신으로 단계가 조정되면
  엑셀을 바꾸고 이 스크립트를 다시 돌린다.

  설계 원칙 — 「신규가 주는 것은 덮어쓰고, 신규에 없는 우리 판단은 이어받는다」
  ──────────────────────────────────────────────────────────────────
  엑셀 기준표 13열은 **발주처·컨설팅이 확정한 사실**이라 그대로 싣는다.
  반면 `수행경로`·`완료판정`·`분류확신도` 같은 열은 **우리가 만든 판단**이고 엑셀에
  없다. 지우면 화면의 «어디서 수행하나»가 통째로 사라지므로, 현행 CSV 에서 **코드로**
  이어받는다(이름은 바뀔 수 있으니 이름으로 이으면 끊긴다). 값이 없는 단계는 비워
  둔다 — 지어내지 않는다.

  엑셀이 확정해 주는 축 — 이름과 뜻
  ────────────────────────────
  · 이행단위   → levelSrc / unitKind   L1 군 · L2 부서 · L3 시설 · L3 공사·용역
  · 주기코드   → cycleCode 명시값. 문자열 파싱보다 **우선**한다.
                 ⚠ GRADE·TERM·MULTIYEAR·BIENNIAL 은 연 단위 고정 회차가 아니다.
  · 이행률 포함 → inRate    **이행률의 분모는 전 단계가 아니라 이 값이 Y 인 단계**다.
  · 이행의무 유형 → dutyKind 정기주기 · 상시·최초 · 조건부(사유 발생) · 조건부(대상 발생)
  · 조문상 의무성 → dutyBasis · 이행 방식 유형 → deliveryKind · EVENT 생성방식 → eventGen
  · 「EVENT생성방식」 시트 → link (연계 시스템·인터페이스·판정 A~D·미확보 항목·조치)

  검증 — 엑셀이 정답지를 갖고 있다 (셋 다 맞아야 파일을 쓴다)
  ────────────────────────────────────────────────────
  ① 「이행률」 시트 ① 이행단위별 대상/이행/미이행 + 합계
  ② 「단계별현황」 시트의 단계별 **판정**(이행·미이행·사유 미발생·대상 미발생) 전건
  ③ 「단계별현황」 시트의 단계별 2025 문서수 = 「문서목록」 집계
  우리가 같은 값을 재현하면 그릇이 맞는 것이다. 어긋나면 파일을 쓰지 않는다.
  (개수를 EXPECT 로 박아 두지 않는다 — 2026-08-18 사고. 엑셀의 자기 집계와 대조한다.)

  실행: python3 tools/build-taxonomy.py --xlsx <경로> [--carry data/…v3.csv]
"""

import argparse
import csv
import io
import json
import os
import re
import sys
from collections import Counter, OrderedDict

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'js', 'doc-taxonomy-data.js')

# 「이행단위」 원문 → (계층, 하위 구분)
UNIT = {
    'L1 군(관리주체)':        ('L1', ''),
    'L2 부서':               ('L2', ''),
    'L3 관리대상(시설)':      ('L3', 'facility'),
    'L3 관리대상(공사·용역)': ('L3', 'work'),
}

# 우리가 만든 판단 — 엑셀에 없어 현행 CSV 에서 코드로 이어받는다
CARRY = ['적용대상', '이행주체', '증빙문서 예시(2025년 실제 문서명)',
         '재난안전과 운영주기(참고)', '수행경로', '완료판정', '분류확신도', '확인필요사유']

# 「단계별현황」 판정 — 문서가 있으면 이행, 없으면 이행의무 유형이 가른다.
# js/cmp-core.js judge() 와 같은 순서여야 한다(화면과 시트가 다른 말을 하면 안 된다).
def judge(n_docs, duty):
    if n_docs > 0:
        return '이행'
    if duty == '조건부(사유 발생)':
        return '사유 미발생'
    if duty == '조건부(대상 발생)':
        return '대상 미발생'
    if duty == '산출불가':
        return '산출불가'
    return '미이행'


def norm(v):
    return '' if v is None else str(v).strip()


def sheet_rows(ws, key='업무단계코드'):
    """머리글 행(첫 칸이 key)을 찾아 그 아래를 dict 로 돌려준다."""
    hdr, out = None, []
    for r in ws.iter_rows(min_row=1, values_only=True):
        if hdr is None:
            if norm(r[0]) == key:
                hdr = [norm(x) for x in r]
            continue
        if norm(r[0]):
            idx = {h: i for i, h in enumerate(hdr) if h}
            out.append({h: norm(r[idx[h]]) for h in idx})
    return out


def load_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    # ── 버전 — 「대시보드」 A2 «기준 v4.0 (88개 업무단계 · 적용 법령 = …)» ──────
    version, scope = '', ''
    if '대시보드' in wb.sheetnames:
        for r in wb['대시보드'].iter_rows(min_row=1, max_row=3, values_only=True):
            m = re.search(r'기준\s*(v[\d.]+)\s*\((.*?)\)', norm(r[0]))
            if m:
                version, scope = m.group(1), m.group(2)
                break

    rows = sheet_rows(wb['기준표'])

    # ── 「EVENT생성방식」 시트 — 연계 개발자 회신이 반영된 **확정 설계** ──────────
    # EVENT 주기 단계가 «언제 어떻게 생기는가»를 정한다. 연계 시스템·인터페이스
    # 번호·미확보 항목·조치까지 갖고 있다. 판정 D(보류)는 전부 계약정보시스템
    # (차세대 e호조) 제공 항목 미수령이다.
    gen = {}
    if 'EVENT생성방식' in wb.sheetnames:
        for g in sheet_rows(wb['EVENT생성방식']):
            gen[g['업무단계코드']] = g

    # 정답지 ① — 「이행률」 시트의 이행단위별 집계
    truth = {}
    seen_hdr = False
    for r in wb['이행률'].iter_rows(min_row=1, values_only=True):
        a = norm(r[0])
        if a == '이행단위':
            seen_hdr = True
            continue
        if seen_hdr and a in UNIT:
            truth[a] = {'total': int(r[1]), 'done': int(r[2]), 'miss': int(r[3])}
        elif seen_hdr and a == '합계':
            truth['합계'] = {'total': int(r[1]), 'done': int(r[2]), 'miss': int(r[3])}
            break

    # 정답지 ② — 「단계별현황」 시트의 단계별 문서수·판정
    status = {}
    if '단계별현황' in wb.sheetnames:
        for s in sheet_rows(wb['단계별현황']):
            status[s['업무단계코드']] = {'docs': int(s.get('2025 문서수') or 0), 'judge': s.get('판정', '')}

    # 문서목록 — 단계명별 건수(정답지 재현에 쓴다)
    ws3 = wb['문서목록']
    h3, docs = None, Counter()
    for r in ws3.iter_rows(min_row=4, values_only=True):
        if h3 is None:
            h3 = [norm(x) for x in r]
            continue
        if r[0] is None and r[1] is None:
            continue
        j = {h: k for k, h in enumerate(h3) if h}
        nm = norm(r[j['업무단계명 (수정 가능)']])
        if nm:
            docs[nm] += 1
    return rows, truth, status, docs, gen, version, scope


def load_carry(path):
    if not path or not os.path.exists(path):
        return {}
    rows = list(csv.DictReader(io.open(path, encoding='utf-8-sig')))
    return {norm(r.get('업무단계코드')): r for r in rows}


def item_of(code):
    """업무단계코드 → 이행항목코드. CIT-01-01 → CIT-01 (앞 두 마디)."""
    p = code.split('-')
    return '-'.join(p[:2]) if len(p) >= 2 else code


def parse_paths(s):
    """'PROGRAM:RSK_REGULAR' · 'ELECTRONIC_DOC' · 'ATTACHMENT' → paths[]"""
    s = norm(s)
    if not s:
        return []
    out = []
    for tok in [t.strip() for t in s.split('|') if t.strip()]:
        kind, _, code = tok.partition(':')
        out.append({'type': kind.strip().upper(), 'code': code.strip()})
    return out


def parse_done(s):
    s = norm(s)
    if not s:
        return {'kind': 'DOC_COUNT', 'key': ''}
    kind, _, key = s.partition(':')
    return {'kind': kind.strip().upper(), 'key': key.strip()}


def build(xrows, carry, gen):
    stages, items = [], OrderedDict()
    for r in xrows:
        code = r['업무단계코드']
        iid = item_of(code)
        old = carry.get(code, {})
        lvl, kind = UNIT.get(r.get('이행단위', ''), ('', ''))

        s = {
            'id': code,
            'itemId': iid,
            'name': r['하위 업무단계명'],
            'law': r['법령근거'],
            'legalCycle': r['법정주기'],
            'opCycle': norm(old.get('재난안전과 운영주기(참고)')),
            'timing': r['수행시점조건'],

            # ── 엑셀이 확정해 준 축 ─────────────────────────────────────────
            'cycleCode': r.get('주기코드', ''),      # 문자열 파싱보다 우선한다
            'levelSrc': lvl,                          # 계층 확정값
            'unitKind': kind,                         # L3 안의 시설 / 공사·용역
            'inRate': (r.get('이행률 포함', '') == 'Y'),   # 이행률 분모 여부
            'dutyKind': r.get('이행의무 유형', ''),
            'dutyBasis': r.get('조문상 의무성', ''),
            'deliveryKind': r.get('이행 방식 유형', ''),
            'eventGen': r.get('EVENT 생성방식', ''),

            # ── 우리 판단 — 코드로 이어받는다(없는 단계는 빈다) ──────────────
            'paths': parse_paths(old.get('수행경로')),
            'doneRule': parse_done(old.get('완료판정')),
            'typeConf': (norm(old.get('분류확신도')) or 'UNKNOWN').upper(),
            'typeNote': norm(old.get('확인필요사유')),
            'target': norm(old.get('적용대상')),
            'actor': norm(old.get('이행주체')),
            'ex': norm(old.get('증빙문서 예시(2025년 실제 문서명)')),
        }
        s['taskType'] = s['paths'][0]['type'] if s['paths'] else 'UNKNOWN'

        # ── EVENT 생성방식 확정 (연계 개발자 회신) ──────────────────────────
        # 값이 없는 단계(EVENT 주기가 아닌 것)는 필드 자체를 만들지 않는다.
        g = gen.get(code)
        if g:
            link = {
                'gen': g.get('최종 생성 방식', ''),
                'sys': g.get('연계 시스템', ''),
                'iface': g.get('연동 인터페이스', ''),
                'grade': g.get('연계 실현 판정', ''),      # A 연계 가능 · B 조건부 가능 · C 연계 불가 · D 원천 미수령
                'missing': g.get('미확보 항목(개발자 회신)', ''),
                'why': g.get('판정 사유', ''),
                'action': g.get('조치', ''),
                'check': g.get('연동 확인 필요', ''),
            }
            if any(link.values()):
                s['link'] = {k: v for k, v in link.items() if v}
        stages.append(s)

        it = items.setdefault(iid, {'id': iid, 'name': r['법정 이행항목명'],
                                    'stageIds': [], 'lawBases': [], 'targets': [],
                                    'actors': [], 'hasCycle': False})
        it['stageIds'].append(code)
        for k, v in (('lawBases', s['law']), ('targets', s['target']), ('actors', s['actor'])):
            if v and v not in it[k]:
                it[k].append(v)
        if s['cycleCode'] != 'EVENT':
            it['hasCycle'] = True
    return stages, list(items.values())


def verify(stages, truth, status, docs):
    """엑셀의 자기 집계 셋을 재현한다 — 못 하면 그릇이 틀린 것이다."""
    errs = []

    # ① 이행률 — 이행단위별
    got = {}
    for s in stages:
        if not s['inRate']:
            continue
        unit = next((k for k, v in UNIT.items() if v == (s['levelSrc'], s['unitKind'])), None)
        if not unit:
            errs.append('이행단위를 못 읽음: ' + s['id'])
            continue
        g = got.setdefault(unit, {'total': 0, 'done': 0, 'miss': 0})
        g['total'] += 1
        if docs.get(s['name'], 0) > 0:
            g['done'] += 1
        else:
            g['miss'] += 1
    tot = {'total': 0, 'done': 0, 'miss': 0}
    for v in got.values():
        for k in tot:
            tot[k] += v[k]
    got['합계'] = tot
    print('\n=== ① 이행률 재현 (엑셀 「이행률」 시트가 정답지) ===')
    for k in list(UNIT) + ['합계']:
        w, g = truth.get(k), got.get(k)
        if not w:
            continue
        ok = g and all(g[x] == w[x] for x in ('total', 'done', 'miss'))
        print('  %-22s 정답 %3d/%3d/%3d · 재현 %s  %s' % (
            k, w['total'], w['done'], w['miss'],
            ('%3d/%3d/%3d' % (g['total'], g['done'], g['miss'])) if g else '(없음)',
            '✓' if ok else '✗'))
        if not ok:
            errs.append('이행률 %s 불일치' % k)

    # ② 판정 · ③ 문서수 — 단계별
    if status:
        print('\n=== ②③ 단계별 판정·문서수 재현 (엑셀 「단계별현황」 시트가 정답지) ===')
        jc, jw = Counter(), Counter()
        for s in stages:
            st = status.get(s['id'])
            if not st:
                errs.append('단계별현황에 없는 단계: ' + s['id'])
                continue
            n = docs.get(s['name'], 0)
            if n != st['docs']:
                errs.append('문서수 불일치 %s (재현 %d · 시트 %d)' % (s['id'], n, st['docs']))
            j = judge(n, s['dutyKind'])
            jw[st['judge']] += 1
            jc[j] += 1
            if j != st['judge']:
                errs.append('판정 불일치 %s (재현 %s · 시트 %s)' % (s['id'], j, st['judge']))
        for k in ['이행', '미이행', '사유 미발생', '대상 미발생', '산출불가']:
            if jw.get(k) or jc.get(k):
                print('  %-10s 시트 %3d · 재현 %3d  %s' % (k, jw.get(k, 0), jc.get(k, 0), '✓' if jw.get(k, 0) == jc.get(k, 0) else '✗'))
        extra = set(status) - {s['id'] for s in stages}
        if extra:
            errs.append('기준표에 없는 단계가 단계별현황에 있음: ' + ', '.join(sorted(extra)))
    return errs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', required=True)
    ap.add_argument('--carry', default=os.path.join(ROOT, 'data', '담양군_안전보건_문서분류기준_v3.csv'))
    ap.add_argument('--version', default='', help='「대시보드」 시트에서 못 읽을 때만 쓴다')
    ap.add_argument('--force', action='store_true', help='검증 실패해도 쓴다(권장하지 않음)')
    a = ap.parse_args()

    xrows, truth, status, docs, gen, version, scope = load_xlsx(a.xlsx)
    version = version or a.version
    if not version:
        sys.exit('분류기준 버전을 「대시보드」 시트에서 읽지 못했습니다 — --version 으로 주세요.')
    carry = load_carry(a.carry)
    stages, items = build(xrows, carry, gen)

    series = Counter(s['id'].split('-')[0] for s in stages)
    print('분류기준 %s — %s' % (version, scope))
    print('이행항목 %d · 업무단계 %d · 계열 %s' % (len(items), len(stages), dict(series)))
    print('  이행단위 :', dict(Counter(s['levelSrc'] + (('/' + s['unitKind']) if s['unitKind'] else '') for s in stages)))
    print('  주기코드 :', dict(Counter(s['cycleCode'] for s in stages).most_common()))
    print('  이행의무 유형 :', dict(Counter(s['dutyKind'] for s in stages).most_common()))
    print('  이행률 포함 Y: %d / %d' % (sum(1 for s in stages if s['inRate']), len(stages)))
    lk = [s for s in stages if s.get('link')]
    print('  EVENT 생성방식: %d단계' % len(lk),
          dict(Counter(s['link'].get('gen', '') for s in lk).most_common()))
    print('  연계 실현 판정 :', dict(Counter(s['link'].get('grade', '—') for s in lk).most_common()))
    print('  우리 판단 이어받음: 수행경로 %d · 완료판정 %d · 적용대상 %d'
          % (sum(1 for s in stages if s['paths']),
             sum(1 for s in stages if s['doneRule']['kind'] == 'PROBE'),
             sum(1 for s in stages if s['target'])))
    if carry:
        dropped = sorted(set(carry) - {s['id'] for s in stages})
        print('  이어받기 CSV 에만 있는 코드(범위 밖으로 빠진 단계) %d' % len(dropped),
              dict(Counter(c.split('-')[0] for c in dropped)))

    errs = verify(stages, truth, status, docs)
    if errs and not a.force:
        print('\n검증 실패 — 파일을 쓰지 않습니다:')
        for e in sorted(set(errs))[:10]:
            print('  ·', e)
        sys.exit(1)

    meta = {'source': os.path.basename(a.xlsx), 'sheet': '기준표', 'version': version,
            'scope': scope,
            'items': len(items), 'stages': len(stages),
            'series': dict(sorted(series.items())),
            'carriedFrom': os.path.basename(a.carry) if carry else '',
            'rateDenominator': sum(1 for s in stages if s['inRate'])}
    hdr = '''/* =========================================================================
 * 업무문서 분류 %s — 법정 이행항목 %d · 하위 업무단계 %d (DYDOCT)
 *   ※ 생성물 — 손으로 고치지 말 것. 원본 엑셀을 고치고 재생성한다.
 *      생성기: tools/build-taxonomy.py
 *      원본:   %s 「기준표」 (%s)
 *
 *   [v3.3 → v4.0] 적용 법령을 중대재해처벌법·산업안전보건법 2종으로 한정
 *   (2026-08-21 군청 미팅). 재난안전법(DSM)·개별 시설법(FAC)·MGT-03~05 가 빠져
 *   95항목 213단계 → %d항목 %d단계. 남은 단계는 코드가 전부 그대로다.
 *
 *   [축] cycleCode(주기 명시값·파싱보다 우선) · levelSrc/unitKind(계층 확정값) ·
 *   inRate(이행률 분모 여부 — 분모는 %d 이지 %d 가 아니다) · dutyKind · dutyBasis ·
 *   deliveryKind · eventGen · link(EVENT 생성방식 — 연계 개발자 회신).
 *
 *   [이어받은 값] paths·doneRule·typeConf·typeNote·target·actor·ex 는 엑셀에 없는
 *   **우리 판단**이라 v3 CSV 에서 코드로 이어받았다. 없는 단계는 비어 있다.
 * ========================================================================= */
''' % (version, len(items), len(stages), os.path.basename(a.xlsx), scope,
       len(items), len(stages), meta['rateDenominator'], len(stages))

    body = 'window.DYDOCT = {\n  META: ' + json.dumps(meta, ensure_ascii=False, indent=2) + ',\n'
    body += '  ITEMS: [\n' + ',\n'.join('    ' + json.dumps(x, ensure_ascii=False) for x in items) + '\n  ],\n'
    body += '  STAGES: [\n' + ',\n'.join('    ' + json.dumps(x, ensure_ascii=False) for x in stages) + '\n  ]\n};\n'
    with open(OUT, 'w', encoding='utf-8') as f:
        f.write(hdr + body)
    print('\n생성: %s (%.1f KB)' % (os.path.relpath(OUT, ROOT), os.path.getsize(OUT) / 1024))


if __name__ == '__main__':
    main()
