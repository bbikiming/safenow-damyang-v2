#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""QA 엑셀 취합 — 테스터에게서 받은 파일들을 한 파일로 모은다.

  실행: python3 tools/collect-qa-xlsx.py <받은 파일 폴더 또는 파일 …> [-o 출력 폴더]
        (출력 폴더를 안 주면 첫 입력 폴더 아래 «취합» 폴더)

[만드는 것]  «QA 담양 중대재해 취합 YYYY-MM-DD.xlsx» + «사진» 폴더
  · 현황          사람 × 메뉴별 성공·실패·오류·보류·확인 불가·문의·남은 것·진행률·사진 수·테스트 환경
  · 확인 필요      실패·오류·문의·보류 줄만 — 메모와 사진 파일 링크를 함께
  · 같은 항목 비교  위험성평가·교육처럼 여럿이 같은 줄을 본 파일 — 한 줄에 사람별 결과를 나란히.
                  여러 사람이 똑같이 실패하면 재현되는 결함이고, 한 사람만 실패하면 환경·조작을 먼저 본다
  · 전체 결과      모든 줄 (필터·피벗용)
  사진은 줄 번호로 찾아 «메뉴_이름_번호_순번.png» 로 꺼낸다(떠 있는 사진 기준).

[테스터 파일을 믿지 않는다]
  · 누구의 무슨 메뉴인지는 숨김 시트 «_정보»에서 읽고, 없으면 파일 이름에서 읽는다
    (메신저로 받으면 파일 이름에 «(1)»이 붙기도 한다).
  · 표 머리 행과 열은 머리 글자로 찾는다 — 테스터가 열이나 행을 끼워 넣어도 틀리지 않는다.
  · 드롭다운을 거치지 않고 붙여 넣은 값(예: «pass»)은 버리지 않고 «알 수 없는 값»으로 드러낸다.
  · 같은 사람·같은 메뉴 파일이 둘이면 나중에 고친 파일을 쓰고 알린다.
  · 엑셀 365 의 «셀에 배치» 사진은 꺼내지 못해 «셀 안 사진 — 원본 파일에서 확인»으로 표시한다.
"""
import os, re, sys, glob, zipfile, html, datetime, unicodedata, argparse

PREFIX = "QA 담양 중대재해"
DEFAULT_RESULTS = ["성공", "실패", "오류", "보류", "확인 불가", "문의"]
ATTENTION = ["오류", "실패", "문의", "보류"]          # 확인 필요 목록 — 앞일수록 먼저 본다
COLOR = {"성공": ("C6EFCE", "006100"), "실패": ("FFC7CE", "9C0006"), "오류": ("F8CBAD", "833C0B"),
         "보류": ("FFEB9C", "7F6000"), "확인 불가": ("D9D9D9", "3A3A3A"), "문의": ("DDEBF7", "1F4E79")}
SKIP_SHEETS = {"안내", "요약", "_정보"}
HEAD_KEYS = {"번호": "no", "누구로": "who", "어디서": "where", "확인할 것": "what",
             "해 볼 것": "steps", "이러면 성공": "expect", "결과": "result", "메모": "memo", "사진": "photo"}

def nfc(s):
    return unicodedata.normalize("NFC", s or "")

def warn(msg, bag):
    bag.append(msg); print("⚠ " + msg)

def sheet_xml_paths(path):
    """시트 이름 → 시트 XML 경로 (셀에 배치된 사진 표시를 찾는 데 쓴다)"""
    out = {}
    with zipfile.ZipFile(path) as z:
        wbx = z.read("xl/workbook.xml").decode("utf-8")
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid_target = {}
    for m in re.finditer(r"<Relationship\b[^>]*>", rels):
        t = m.group(0)
        i, g = re.search(r'Id="([^"]+)"', t), re.search(r'Target="([^"]+)"', t)
        if i and g:
            tg = g.group(1).lstrip("/")
            rid_target[i.group(1)] = tg if tg.startswith("xl/") else "xl/" + tg
    for m in re.finditer(r"<sheet\b[^>]*>", wbx):
        t = m.group(0)
        n, r = re.search(r'name="([^"]*)"', t), re.search(r'r:id="([^"]+)"', t)
        if n and r and r.group(1) in rid_target:
            out[html.unescape(n.group(1))] = rid_target[r.group(1)]
    return out

def incell_image_rows(path, sheet_xml, col_letter):
    """엑셀 365 «셀에 배치» 사진은 셀 값(vm 속성)으로 들어간다 — 그 줄 번호만 알려 준다"""
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read(sheet_xml).decode("utf-8")
    except KeyError:
        return set()
    rows = set()
    for m in re.finditer(r'<c r="([A-Z]+)(\d+)"[^>]*\bvm="\d+"', xml):
        if m.group(1) == col_letter:
            rows.add(int(m.group(2)))
    return rows

def read_file(path, photo_dir, warns):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    wb = load_workbook(path, data_only=True)
    base = nfc(os.path.basename(path))
    info = {}
    if "_정보" in wb.sheetnames:
        for a, b in wb["_정보"].iter_rows(min_row=1, max_col=2, values_only=True):
            if a:
                info[str(a)] = b
    menu, person = info.get("메뉴"), info.get("담당자")
    if not (menu and person):
        m = re.match(r"^%s (.+) (\S+?)(?: ?\(\d+\))?\.xlsx$" % re.escape(PREFIX), base)
        if not m:
            warn("누구의 무슨 파일인지 알 수 없어 건너뜀 — %s" % base, warns)
            return None
        menu, person = m.group(1), m.group(2)
        warn("«_정보» 시트가 없어 파일 이름으로 읽음 — %s" % base, warns)
    results = [x for x in str(info.get("결과 값") or "").split(",") if x] or DEFAULT_RESULTS
    env = {}
    if "안내" in wb.sheetnames:
        for row in wb["안내"].iter_rows(min_row=1, max_row=40, max_col=2, values_only=True):
            if row[0] in ("브라우저", "기기", "시작한 날", "테스트 주소"):
                env[row[0]] = row[1]
    xml_of = sheet_xml_paths(path)
    rows, photos = [], 0
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS or ws.sheet_state != "visible":
            continue
        head_row, cols = None, {}
        for r in range(1, 12):
            if str(ws.cell(row=r, column=1).value or "").strip() == "번호":
                head_row = r; break
        if not head_row:
            warn("%s / %s — 표 머리(번호)를 찾지 못해 건너뜀" % (base, ws.title), warns)
            continue
        for c in range(1, ws.max_column + 1):
            h = str(ws.cell(row=head_row, column=c).value or "").strip()
            for k, key in HEAD_KEYS.items():
                if h.startswith(k) and key not in cols:
                    cols[key] = c
        if "result" not in cols:
            warn("%s / %s — «결과» 열을 찾지 못해 건너뜀" % (base, ws.title), warns)
            continue
        by_row = {}
        for r in range(head_row + 1, ws.max_row + 1):
            no = ws.cell(row=r, column=cols.get("no", 1)).value
            try:
                no = int(no)
            except (TypeError, ValueError):
                continue
            g = lambda key: ws.cell(row=r, column=cols[key]).value if key in cols else None
            res = str(g("result") or "").strip()
            rec = {"menu": menu, "person": person, "kind": info.get("종류") or "", "code": info.get("파일 코드") or "",
                   "sheet": ws.title, "no": no, "who": g("who") or "",
                   "where": g("where") or "", "what": g("what") or "", "result": res,
                   "unknown": bool(res) and res not in results, "memo": g("memo") or "",
                   "photos": [], "incell": False, "file": base, "row": r}
            rows.append(rec); by_row[r] = rec
        if not by_row:
            continue
        data_rows = sorted(by_row)
        def owner(r1):
            # 사진은 줄 칸에 딱 맞지 않고 걸쳐 있기 마련이다 — 사진 윗변이 걸린 줄(없으면 바로 윗줄)
            if r1 in by_row:
                return by_row[r1]
            above = [x for x in data_rows if x <= r1]
            return by_row[above[-1]] if above else None
        for k, img in enumerate(getattr(ws, "_images", [])):
            anc = getattr(img, "anchor", None)
            if isinstance(anc, str):                      # 저장 전 개체는 «J5» 같은 문자열이다
                mm = re.match(r"^[A-Z]+(\d+)$", anc)
                top = int(mm.group(1)) if mm else None
            else:
                frm = getattr(anc, "_from", None)
                top = frm.row + 1 if frm is not None else None
            if top is None:
                continue
            rec = owner(top)
            if not rec:
                continue
            ext = (getattr(img, "format", None) or "png").lower().replace("jpeg", "jpg")
            name = "%s_%s_%03d_%d.%s" % (menu, person, rec["no"], len(rec["photos"]) + 1, ext)
            name = re.sub(r"[\\/:*?\"<>|]", "_", name)
            try:
                data = img._data()
            except Exception:
                warn("%s / %s %d번 — 사진을 읽지 못함" % (base, ws.title, rec["no"]), warns)
                continue
            with open(os.path.join(photo_dir, name), "wb") as fh:
                fh.write(data)
            rec["photos"].append(name); photos += 1
        if "photo" in cols and ws.title in xml_of:
            for r in incell_image_rows(path, xml_of[ws.title], get_column_letter(cols["photo"])):
                if r in by_row:
                    by_row[r]["incell"] = True
    return {"menu": menu, "person": person, "kind": info.get("종류") or "", "results": results,
            "rows": rows, "photos": photos, "env": env, "file": base, "path": path}

def main():
    ap = argparse.ArgumentParser(description="QA 엑셀 취합")
    ap.add_argument("inputs", nargs="*", help="받은 파일 폴더 또는 파일")
    ap.add_argument("-o", "--out", help="출력 폴더 (기본: 첫 입력 폴더/취합)")
    a = ap.parse_args()
    root = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
    inputs = a.inputs or [os.path.join(root, "docs", "planning", "QA엑셀-v1")]
    if not a.inputs:
        print("(입력 폴더를 주지 않아 배포 폴더를 읽습니다 — 받은 파일은 폴더를 지정하세요)")
    first = inputs[0] if os.path.isdir(inputs[0]) else os.path.dirname(inputs[0])
    out_dir = os.path.abspath(a.out or os.path.join(first, "취합"))
    photo_dir = os.path.join(out_dir, "사진")
    os.makedirs(photo_dir, exist_ok=True)
    for old in glob.glob(os.path.join(photo_dir, "*")):
        os.remove(old)                     # 지난 취합의 사진이 섞이지 않게 — 이 폴더는 이 스크립트 것이다

    paths = []
    for p in inputs:
        if os.path.isdir(p):
            for dp, _, fs in os.walk(p):
                if os.path.abspath(dp).startswith(out_dir):
                    continue
                for fn in fs:
                    n = nfc(fn)
                    if n.endswith(".xlsx") and not n.startswith("~$") and "취합" not in n:
                        paths.append(os.path.join(dp, fn))
        elif p.endswith(".xlsx"):
            paths.append(p)
    if not paths:
        print("✖ 읽을 엑셀 파일이 없다: %s" % ", ".join(inputs)); sys.exit(1)

    warns, files = [], {}
    for p in sorted(paths):
        f = read_file(p, photo_dir, warns)
        if not f:
            continue
        key = (f["menu"], f["person"])
        if key in files:
            keep = max(files[key], f, key=lambda x: os.path.getmtime(x["path"]))
            warn("%s · %s 파일이 둘이다 — 나중에 고친 «%s» 를 씀" % (key[0], key[1], keep["file"]), warns)
            files[key] = keep
        else:
            files[key] = f
    rows = [r for f in files.values() for r in f["rows"]]
    results = next(iter(files.values()))["results"] if files else DEFAULT_RESULTS

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.formatting.rule import CellIsRule, DataBarRule
    from openpyxl.utils import get_column_letter as L
    FONT = "맑은 고딕"
    fb, fbb = Font(name=FONT, size=10.5), Font(name=FONT, size=10.5, bold=True)
    fh = Font(name=FONT, size=10.5, bold=True, color="FFFFFF")
    fill_h = PatternFill("solid", fgColor="1F4E79")
    thin = Side(style="thin", color="BFBFBF"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top"); ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()
    CENTER = {"담당자", "번호", "결과", "사람 수", "브라우저", "기기", "시작한 날"} | set(results)
    def table(ws, heads, widths, data, title, note=None, center=()):
        center = set(center) | CENTER
        ws.cell(row=1, column=1, value=title).font = Font(name=FONT, size=15, bold=True)
        if note:
            ws.cell(row=2, column=1, value=note).font = Font(name=FONT, size=10.5, color="7F7F7F", italic=True)
        for k, (h, w) in enumerate(zip(heads, widths), 1):
            c = ws.cell(row=4, column=k, value=h); c.font = fh; c.fill = fill_h; c.alignment = ctr; c.border = bd
            ws.column_dimensions[L(k)].width = w
        for i, row in enumerate(data):
            for k, v in enumerate(row, 1):
                c = ws.cell(row=5 + i, column=k, value=v); c.font = fb; c.border = bd
                c.alignment = ctr if isinstance(v, (int, float)) or heads[k - 1] in center else wrap
        last = 4 + max(1, len(data))
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = "A4:%s%d" % (L(len(heads)), last)
        return last
    def color_results(ws, rng):
        for name in results:
            if name in COLOR:
                bg, fg = COLOR[name]
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"%s"' % name],
                    fill=PatternFill("solid", fgColor=bg), font=Font(name=FONT, size=10.5, bold=True, color=fg)))

    # 현황
    ws = wb.active; ws.title = "현황"
    heads = ["담당자", "메뉴", "전체"] + results + ["남은 것", "진행률", "사진", "알 수 없는 값", "브라우저", "기기", "시작한 날", "파일"]
    data = []
    for (menu, person), f in sorted(files.items(), key=lambda kv: (kv[0][1], kv[0][0])):
        rs = f["rows"]; n = len(rs)
        cnt = [sum(1 for r in rs if r["result"] == x) for x in results]
        left = sum(1 for r in rs if not r["result"])
        data.append([person, menu, n] + cnt + [left, (n - left) / n if n else 0, f["photos"],
                     sum(1 for r in rs if r["unknown"]), f["env"].get("브라우저"), f["env"].get("기기"),
                     f["env"].get("시작한 날"), f["file"]])
    last = table(ws, heads, [14, 16, 7] + [8] * len(results) + [8, 10, 7, 9, 10, 10, 12, 44], data,
                 "QA 취합 — 현황", "파일 %d개 · 줄 %d개 · 사진 %d장 · 취합 %s" % (
                     len(files), len(rows), sum(f["photos"] for f in files.values()), datetime.date.today().isoformat()))
    pc = 4 + len(results) + 1
    for r in range(5, last + 1):
        ws.cell(row=r, column=pc).number_format = "0%"
    ws.conditional_formatting.add("%s5:%s%d" % (L(pc), L(pc), last), DataBarRule(
        start_type="num", start_value=0, end_type="num", end_value=1, color="63BE7B", showValue=True))
    for x in ("실패", "오류"):
        if x in results:
            c = 4 + results.index(x)
            ws.conditional_formatting.add("%s5:%s%d" % (L(c), L(c), last), CellIsRule(
                operator="greaterThan", formula=["0"], font=Font(name=FONT, size=10.5, bold=True, color="C00000")))
    # 메뉴별 합계 — 같은 메뉴를 여럿이 본 경우(위험성평가·교육) 한 줄로
    r0 = last + 3
    ws.cell(row=r0 - 1, column=1, value="메뉴별 합계").font = Font(name=FONT, size=12, bold=True, color="1F4E79")
    for k, h in enumerate(["메뉴", "사람 수", "전체"] + results + ["남은 것"], 1):
        c = ws.cell(row=r0, column=k, value=h); c.font = fh; c.fill = fill_h; c.alignment = ctr; c.border = bd
    # 여럿이 보는 파일(상세)을 앞에, 나머지는 이름 순
    kind_of = {m: f["kind"] for (m, _), f in files.items()}
    menus = sorted({m for m, _ in files}, key=lambda m: (0 if kind_of.get(m) == "상세" else 1, m))
    for i, m in enumerate(menus, 1):
        rs = [r for r in rows if r["menu"] == m]
        vals = [m, len({r["person"] for r in rs}), len(rs)] + [sum(1 for r in rs if r["result"] == x) for x in results] + \
               [sum(1 for r in rs if not r["result"])]
        for k, v in enumerate(vals, 1):
            c = ws.cell(row=r0 + i, column=k, value=v); c.font = fb; c.border = bd; c.alignment = ctr if k > 1 else wrap

    # 확인 필요
    ws2 = wb.create_sheet("확인 필요")
    order = {x: i for i, x in enumerate(ATTENTION)}
    need = [r for r in rows if r["result"] in ATTENTION or r["unknown"] or r["incell"]]
    need.sort(key=lambda r: (order.get(r["result"], -1 if r["unknown"] else 9), r["menu"], r["no"], r["person"]))
    data2 = []
    for r in need:
        ph = ", ".join(r["photos"]) if r["photos"] else ("셀 안 사진 — 원본 파일에서 확인" if r["incell"] else "")
        data2.append([r["menu"], r["sheet"], r["no"], r["what"],
                      r["result"] + (" (알 수 없는 값)" if r["unknown"] else ""), r["person"], r["memo"], ph, r["file"]])
    last2 = table(ws2, ["메뉴", "시트", "번호", "확인할 것", "결과", "담당자", "메모", "사진", "파일"],
                  [14, 18, 6, 30, 12, 9, 44, 30, 40], data2, "확인 필요 — 실패·오류·문의·보류",
                  "오류 → 실패 → 문의 → 보류 순. 사진 칸을 누르면 «사진» 폴더의 파일이 열립니다.")
    for i, r in enumerate(need):
        if r["photos"]:
            c = ws2.cell(row=5 + i, column=8)
            c.hyperlink = "사진/" + r["photos"][0]
            c.font = Font(name=FONT, size=10.5, color="0563C1", underline="single")
    color_results(ws2, "E5:E%d" % last2)

    # 같은 항목 비교 — 여럿이 같은 줄을 본 파일만
    ws3 = wb.create_sheet("같은 항목 비교")
    multi = sorted({r["menu"] for r in rows if r["kind"] == "상세"} |
                   {m for m in {r["menu"] for r in rows} if len({r["person"] for r in rows if r["menu"] == m}) > 1})
    people = sorted({r["person"] for r in rows if r["menu"] in multi})
    data3 = []
    for m in multi:
        items = {}
        for r in rows:
            if r["menu"] == m:
                items.setdefault(r["no"], {"what": r["what"], "sheet": r["sheet"], "by": {}})
                items[r["no"]]["by"][r["person"]] = r["result"]
        for no in sorted(items):
            it = items[no]
            res = [it["by"].get(p, "") for p in people]
            data3.append([m, it["sheet"], no, it["what"]] + res +
                         [sum(1 for x in res if x in ("실패", "오류")), sum(1 for x in res if x == "성공")])
    last3 = table(ws3, ["메뉴", "시트", "번호", "확인할 것"] + people + ["실패·오류", "성공"],
                  [14, 18, 6, 32] + [9] * len(people) + [9, 7], data3, "같은 항목 비교", center=people,
                  note="같은 줄을 여러 사람이 봤을 때 — 여럿이 실패하면 재현되는 결함, 한 사람만 실패하면 그 사람의 환경·조작을 먼저 봅니다.")
    if people:
        color_results(ws3, "%s5:%s%d" % (L(5), L(4 + len(people)), last3))

    # 전체 결과
    ws4 = wb.create_sheet("전체 결과")
    data4 = [[r["menu"], r["sheet"], r["no"], r["who"], r["where"], r["what"], r["result"],
              r["person"], r["memo"], ", ".join(r["photos"]), r["file"]] for r in rows]
    last4 = table(ws4, ["메뉴", "시트", "번호", "누구로", "어디서", "확인할 것", "결과", "담당자", "메모", "사진", "파일"],
                  [14, 18, 6, 12, 22, 30, 10, 9, 40, 24, 40], data4, "전체 결과")
    color_results(ws4, "G5:G%d" % last4)

    out = os.path.join(out_dir, "%s 취합 %s.xlsx" % (PREFIX, datetime.date.today().isoformat()))
    wb.save(out)
    print("\n✔ 취합 — 파일 %d개 · 줄 %d개 · 확인 필요 %d줄 · 사진 %d장" % (
        len(files), len(rows), len(need), sum(f["photos"] for f in files.values())))
    print("  → %s" % out)
    print("  → %s" % photo_dir)
    if warns:
        print("  ⚠ 알림 %d건 (위 참고)" % len(warns))

if __name__ == "__main__":
    main()
