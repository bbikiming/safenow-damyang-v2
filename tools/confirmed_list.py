# -*- coding: utf-8 -*-
"""
확정문서목록 읽기 — build-ledger-2025.py · build-confirmed-docs.py 공용

  원본: 담양군_안전보건_확정문서목록_연도별_<날짜>.xlsx (㈜다온플레이스, 2026-09-18)
    · 「요약」  : 연도별 수록 현황 표(기준연도 · 문서×단계 행 · 문서 · 부서 · 연결 단계)
    · 연도 시트: 2021하반기 · 2022 · 2023 · 2024 · 2025 · 2026상반기
                 1행 = 문서 × 업무단계. 한 문서가 여러 단계면 단계마다 행이 생긴다.

  문서의 정체 = 부서 + 보고일자 + 문서명 + 수발신자
  ───────────────────────────────────────────
  「요약」의 «문서» 수가 이 키의 고유 개수와 정확히 같다(연도 6개 전건). 그래서
  이 키를 문서 하나로 본다. 같은 키·같은 단계가 두 번 나오는 행(2025 15쌍 등)은
  같은 증빙을 두 번 센 것이라 하나로 친다 — 한 문서가 한 단계를 두 번 증빙하지
  않는다.

  검증 — 「요약」 표를 재현하지 못하면 멈춘다
  ─────────────────────────────────────
  행 수 · 문서 수 · 부서 수 셋을 연도마다 대조한다. 고정 건수를 검증값으로 박지
  않는다(§5 2026-08-18 EXPECT 사고) — 원본이 스스로 말하는 수와 맞추는 것이다.

  ⚠ 「연결 단계」 열은 재현되지 않는다(2026-09-28 실측) — 2021 하반기만 맞고 나머지는
  고유 단계코드 수보다 1~13 작다(2025: 요약 56 · 코드 69). 이행률 포함(Y)만·주기
  있는 것만·2건 이상만 등 여섯 가지 정의를 대 봤지만 어느 것도 전 연도에 맞지 않아
  정의를 알 수 없다. 그래서 **멈추지 않고 보고만 한다** — 우리가 정의를 지어내
  맞추면 그게 날조다. 화면은 이 열을 쓰지 않고 단계코드에서 직접 센다.
  작성처(㈜다온플레이스 분류 담당)에 정의를 확인할 항목이다.
"""

import re
import sys
from collections import OrderedDict

KEY = ('부서', '보고일자', '문서명', '수발신자')
DIR_OF = {'접수(외부)': 'in', '생산': 'out', '접수(내부)': 'internal'}


def norm(v):
    return '' if v is None else str(v).strip()


def norm_date(s):
    """'2025.01.02' → '2025-01-02'. 판독 불가는 원문 유지."""
    s = norm(s)
    m = re.match(r'^(\d{4})[.\-/](\d{1,2})[.\-/](\d{1,2})\.?$', s)
    return '%s-%02d-%02d' % (m.group(1), int(m.group(2)), int(m.group(3))) if m else s


def key_of(r):
    return tuple(r.get(k, '') for k in KEY)


def sheet_period(title):
    """'2021하반기' → (2021, 'H2', '2021년 하반기') · '2025' → (2025, '', '2025년')"""
    m = re.match(r'^(\d{4})(상반기|하반기)?$', title)
    if not m:
        return None
    y, p = int(m.group(1)), m.group(2) or ''
    return y, {'상반기': 'H1', '하반기': 'H2', '': ''}[p], '%d년%s' % (y, (' ' + p) if p else '')


def load(path):
    """→ OrderedDict(year → {period, label, rows[], docs OrderedDict(key → {row, stages[]})}), summary{label→dict}"""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    years = OrderedDict()
    for ws in wb.worksheets:
        sp = sheet_period(ws.title)
        if not sp:
            continue
        y, per, label = sp
        it = ws.iter_rows(values_only=True)
        hdr = [norm(x) for x in next(it)]
        rows = []
        for r in it:
            if r[1] is None and r[3] is None:
                continue
            rows.append({h: norm(r[i]) for i, h in enumerate(hdr) if h})
        docs = OrderedDict()
        for r in rows:
            k = key_of(r)
            d = docs.get(k)
            if d is None:
                d = docs[k] = {'row': r, 'stages': []}
            if r['단계코드'] and r['단계코드'] not in d['stages']:
                d['stages'].append(r['단계코드'])
        years[y] = {'period': per, 'label': label, 'sheet': ws.title, 'rows': rows, 'docs': docs}

    # 「요약」 연도별 수록 현황 표
    summary = {}
    ws = wb['요약']
    on = False
    for r in ws.iter_rows(values_only=True):
        c = [norm(x) for x in r]
        if c and c[0] == '기준연도':
            on = True
            continue
        if on:
            if not c or not c[0] or c[0] == '계':
                if c and c[0] == '계':
                    break
                continue
            summary[c[0]] = {'rows': int(c[1]), 'docs': int(c[2]), 'depts': int(c[3]), 'stages': int(c[4])}
    return years, summary


def verify(years, summary, stage_ids):
    """「요약」 재현 + 단계코드 실재 → (멈출 이유[], 보고만 할 차이[]). 빈 목록 = 통과."""
    bad, warn = [], []
    if not summary:
        bad.append('「요약」 연도별 수록 현황 표를 찾지 못했습니다')
    for y, Y in years.items():
        s = summary.get(Y['label'])
        if not s:
            bad.append('%s — 「요약」에 이 연도 행이 없습니다' % Y['label'])
            continue
        got = {
            'rows': len(Y['rows']),
            'docs': len(Y['docs']),
            'depts': len({k[0] for k in Y['docs']}),
            'stages': len({c for d in Y['docs'].values() for c in d['stages']}),
        }
        for f in ('rows', 'docs', 'depts'):
            if got[f] != s[f]:
                bad.append('%s %s — 요약 %d · 재현 %d' % (Y['label'], f, s[f], got[f]))
        if got['stages'] != s['stages']:
            warn.append('%s 연결 단계 — 요약 %d · 고유 단계코드 %d' % (Y['label'], s['stages'], got['stages']))
        miss = sorted({c for d in Y['docs'].values() for c in d['stages']} - set(stage_ids))
        if miss:
            bad.append('%s — 분류기준에 없는 단계코드 %d종: %s' % (Y['label'], len(miss), ', '.join(miss[:5])))
    return bad, warn


def report(warn):
    if warn:
        print('  (보고) 「요약」 연결 단계 열은 정의를 알 수 없어 대조만 합니다:')
        for w in warn:
            print('    ·', w)


def fail(bad, what):
    print('\n%s 검증 실패 — 파일을 쓰지 않습니다:' % what)
    for b in bad:
        print('  ·', b)
    sys.exit(1)
