#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""QA 엑셀 배포본 생성기 — 원본 문서 하나 → 사람별 QA 파일.

  실행: python3 tools/build-qa-xlsx.py            (만들기)
        python3 tools/build-qa-xlsx.py --check    (쓰지 않고 원본만 검사)

[왜 있는가]
QA 인원이 6명이고 위험성평가·안전보건교육은 **전원이 같은 줄**을 각자 파일로 받는다.
손으로 12개를 만들면 한 곳을 고칠 때 12개를 다 고쳐야 하고 반드시 갈린다. 줄은
원본 문서(docs/planning/검수-QA엑셀-v1.md) 한 곳에만 두고 파일은 여기서 만든다.
**엑셀을 손으로 고치지 말 것** — 원본을 고치고 다시 돌린다.

[테스터가 보는 모양 — 2026-09-28 v1.1 단순화]
  한 줄 = 한 기능. 번호 · 누구로 · 어디서 · 확인할 것 · 해 볼 것 · 이러면 성공
  (중요도 칸은 없다 — 2026-09-28 사용자: «중요도는 필요 없다». 모든 줄을 다 본다)
  → 테스터는 «결과»(드롭다운) · «메모» · «사진»(캡처 붙여넣기) 세 칸만 채운다.
  · 결과 값·색은 원본 «결과 고르는 법» 표에서 읽는다 — 드롭다운과 안내가 갈리지 않는다.
  · 시트 보호를 걸지 않는다 — 보호된 시트에서는 캡처 붙여넣기·열 너비 조절이 막혀
    «셀이 보호되어 있습니다» 창이 뜬다. 테스터가 막히는 지점을 하나라도 줄인다.
  · 결함 기록 시트를 따로 두지 않는다 — 실패한 줄에 메모·사진을 바로 남기고,
    모으는 일은 취합 스크립트(tools/collect-qa-xlsx.py)가 한다.

[사람마다 다른 값 — 2026-09-29 서버 QA]
  서버에서는 사람마다 계정·부서가 다르다. 원본 줄에는 치환어를 쓰고 파일마다 채운다.
  · 박안전N · 하정수N · {내 부서} · {이름} — 배정 표의 «재난안전과 계정 · 부서 계정 · 내 부서» 칸
  · {진행자} · {진행자 부서} · {대상 부서} — 기본값 «정기평가 진행자 · 진행자 몫 부서»
  · «누구로»가 «진행자: …» 인 줄은 진행자 파일에만 들어간다(한 해 1건뿐인 정기평가를 6명이
    함께 진행하기 때문 — 만들기·전달·상신은 한 사람만 한다). 다른 사람 파일은 번호가 건너뛴다.
  · 파일 절 안의 «### 안내: 제목» 은 공통 안내에서 첫 낱말이 같은 절을 그 파일만 바꾼다
    (첫 낱말이 없는 절이면 새로 더한다). 서버용 파일과 프로토타입용 파일이 한 원본에 섞여 있을 수 있다.
  · 치환어 바로 뒤에 조사를 붙이지 않는다 — 값마다 받침이 달라(담양읍을 · 건설과를) 틀린 조사가 된다.

[멈추는 조건] 원본에 형식 오류가 있으면 **파일을 한 개도 쓰지 않는다** — 칸 수가 6이 아니거나,
번호가 중복·접두어 불일치이거나, «누구로»가 계정 표에 없거나, «어디서»의 메뉴 경로가
실제 메뉴(js/layout.js NAV)와 다르거나, «N번» 참조가 없는 줄을 가리키거나,
배정 표가 스모크 파일을 빠뜨리거나 두 번 주면, 치환어를 채울 값이 없거나 치환어 뒤에 조사가
붙어 있으면, 진행자 줄을 빼고 나니 누군가에게 빈 시트·없는 «N번»이 생기면.
"""
import io, os, re, sys, datetime, unicodedata

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
SRC = os.path.join(ROOT, "docs", "planning", "검수-QA엑셀-v1.md")
OUT = os.path.join(ROOT, "docs", "planning", "QA엑셀-v1")
NAV_JS = os.path.join(ROOT, "js", "layout.js")
CHECK = "--check" in sys.argv[1:]
PREFIX = "QA 담양 중대재해"
FORMAT = "QA엑셀 v1.1"
BROWSERS = ["크롬", "엣지", "웨일", "사파리", "파이어폭스", "기타"]
DEVICES = ["윈도우 PC", "맥", "태블릿", "휴대폰"]
RESULTS_SECTION = "결과 고르는 법"
ACCOUNTS_SECTION = "누구로"          # 섹션 제목에 이 말이 들어간 절의 표
TC_HEAD_SRC = ["번호", "누구로", "어디서", "확인할 것", "해 볼 것", "이러면 성공"]

def die(msg):
    print("✖ " + msg); sys.exit(1)

def nfc(s):
    return unicodedata.normalize("NFC", s)

# ── ① 원본 파싱 ────────────────────────────────────────────────────────────
lines = io.open(SRC, encoding="utf-8").read().split("\n")

def cells(l):
    return [c.strip() for c in l.strip().strip("|").split("|")]

def is_sep(l):
    return bool(re.match(r"^\|[\s\-:|]+\|\s*$", l.strip()))

meta, assign, guide, files = {}, [], [], []
assign_head, assign_info = None, {}
sec = cur_file = cur_sheet = cur_guide = cur_fguide = None
RULE = re.compile(r"^-{3,}$")          # 절 사이 가로줄 — 안내 문단으로 들어가면 «• ---»로 찍힌다
i = 0
while i < len(lines):
    l = lines[i]
    m2 = re.match(r"^## (.+)$", l)
    m3 = re.match(r"^### (.+)$", l)
    if m2:
        title = m2.group(1).strip()
        cur_sheet = cur_guide = cur_fguide = None
        mf = re.match(r"^파일:\s*(.+?)\s*\((\w+)\s*·\s*(상세|스모크)\)$", title)
        sec = {"기본값": "meta", "배정": "assign", "안내": "guide"}.get(title)
        if mf:
            sec = "file"
            cur_file = {"name": mf.group(1), "code": mf.group(2), "kind": mf.group(3), "desc": [], "sheets": [], "guide": []}
            files.append(cur_file)
        i += 1; continue
    if m3 and sec == "guide":
        cur_guide = (m3.group(1).strip(), []); guide.append(cur_guide); i += 1; continue
    if m3 and sec == "file":
        mg = re.match(r"^안내:\s*(.+)$", m3.group(1).strip())
        if mg:
            cur_sheet = None
            cur_fguide = (mg.group(1).strip(), []); cur_file["guide"].append(cur_fguide); i += 1; continue
        ms = re.match(r"^시트:\s*(.+)$", m3.group(1).strip())
        if not ms:
            die("%d행 — 파일 절 안의 ### 는 «시트: 이름» 또는 «안내: 제목» 이어야 한다" % (i + 1))
        cur_fguide = None
        cur_sheet = {"name": ms.group(1).strip(), "desc": [], "rows": []}
        cur_file["sheets"].append(cur_sheet); i += 1; continue

    if sec == "meta" and l.startswith("|") and not is_sep(l):
        c = cells(l)
        if len(c) >= 2 and c[0] != "항목":
            meta[c[0]] = c[1]
    elif sec == "assign" and l.startswith("|") and not is_sep(l):
        c = cells(l)
        if c[0] == "이름":
            assign_head = c
        else:
            assign.append((c[0], [x.strip() for x in c[1].split(",") if x.strip()]))
            assign_info[c[0]] = dict(zip(assign_head or [], c))
    elif sec == "guide" and cur_guide is not None:
        if l.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].startswith("|"):
                if not is_sep(lines[i]):
                    tbl.append(cells(lines[i]))
                i += 1
            cur_guide[1].append(("table", tbl)); continue
        if l.strip() and not RULE.match(l.strip()):
            cur_guide[1].append(("p", l.strip()))
    elif sec == "file" and cur_fguide is not None:
        if l.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].startswith("|"):
                if not is_sep(lines[i]):
                    tbl.append(cells(lines[i]))
                i += 1
            cur_fguide[1].append(("table", tbl)); continue
        if l.strip() and not RULE.match(l.strip()):
            cur_fguide[1].append(("p", l.strip()))
    elif sec == "file":
        if l.startswith(">"):
            (cur_sheet if cur_sheet is not None else cur_file)["desc"].append(l.lstrip("> ").strip())
        elif l.startswith("|") and cur_sheet is not None and not is_sep(l):
            c = cells(l)
            if c[0] != "번호":
                if len(c) != len(TC_HEAD_SRC):
                    die("%d행 — 칸이 %d개다(%d이어야 한다). 셀 안에 세로 막대가 있으면 이렇게 된다: %s" % (i + 1, len(c), len(TC_HEAD_SRC), l[:80]))
                cur_sheet["rows"].append(c)
            elif c != TC_HEAD_SRC:
                die("%d행 — 표 머리가 %s 여야 한다" % (i + 1, " · ".join(TC_HEAD_SRC)))
    i += 1

def guide_table(title_part):
    for t, blocks in guide:
        if title_part in t:
            for k, v in blocks:
                if k == "table":
                    return v
    return None

# 결과 값·색 — 드롭다운과 안내 표가 같은 값을 쓰도록 여기서 읽는다
COLOR = {"초록": ("C6EFCE", "006100"), "빨강": ("FFC7CE", "9C0006"), "주황": ("F8CBAD", "833C0B"),
         "노랑": ("FFEB9C", "7F6000"), "회색": ("D9D9D9", "3A3A3A"), "파랑": ("DDEBF7", "1F4E79")}
rt = guide_table(RESULTS_SECTION)
if not rt or rt[0][:2] != ["결과", "이럴 때 고르세요"]:
    die("안내의 «%s» 표(결과 · 이럴 때 고르세요 · 색)를 찾지 못했다" % RESULTS_SECTION)
RESULTS = []
for r in rt[1:]:
    if len(r) < 3 or r[2] not in COLOR:
        die("결과 «%s» 의 색이 %s 중 하나가 아니다" % (r[0], "·".join(COLOR)))
    if "," in r[0]:
        die("결과 값에 쉼표를 쓸 수 없다(드롭다운 목록이 갈린다): %s" % r[0])
    RESULTS.append((r[0], r[1], r[2]))
RESULT_NAMES = [r[0] for r in RESULTS]
FAIL_LIKE = [x for x in RESULT_NAMES if x in ("실패", "오류", "문의")]

at = guide_table(ACCOUNTS_SECTION)
if not at or at[0][0] != "이름":
    die("안내의 «%s» 절에서 계정 표(이름 · 누구 · 주로 하는 것)를 찾지 못했다" % ACCOUNTS_SECTION)
ACCOUNTS = [r[0] for r in at[1:]]

# ── 사람마다 다른 값 ─────────────────────────────────────────────────────────
LEAD = meta.get("정기평가 진행자", "").strip()
LEAD_DEPT = meta.get("진행자 몫 부서", "").strip()
LEAD_TAG = "진행자:"
ROLE_OF = {"박안전N": "재난안전과", "하정수N": "담당 공무원", "QA시스템관리": "시스템 관리"}     # «누구로» 칸에 계정과 함께 보이는 권한
TOKEN_RE = re.compile(r"\{[^{}]+\}|박안전N|하정수N")
# 치환어 바로 뒤의 조사 — 값마다 받침이 달라 틀린 조사가 된다(담양읍를 · 하정수3로). 에·의·에서·도·만은 괜찮다.
JOSA_RE = re.compile(r"(\{[^{}]+\}|박안전N|하정수N)(으로|로|은|는|이|가|을|를|와|과|이나|나|이다|이고|이라)(?![가-힣])")

def no_detail(person):
    """배정 표 «상세 파일» 칸이 «제외»면 위험성평가·교육 상세 파일을 받지 않는다(둘러보기 전용 인원)"""
    return assign_info.get(person, {}).get("상세 파일", "").strip() == "제외"

def tokens_for(person):
    info = assign_info.get(person, {})
    t = {"{이름}": person}
    for tok, col in (("박안전N", "재난안전과 계정"), ("하정수N", "부서 계정"), ("{내 부서}", "내 부서")):
        if info.get(col):
            t[tok] = info[col]
    if LEAD:
        t["{진행자}"] = LEAD
    if LEAD_DEPT:
        t["{진행자 부서}"] = LEAD_DEPT
        depts = [assign_info.get(n, {}).get("내 부서", "") for n, _ in assign if not no_detail(n)]
        if all(depts):
            t["{대상 부서}"] = " · ".join(depts + [LEAD_DEPT])
    return t

def subst(text, person):
    t = tokens_for(person)
    out = TOKEN_RE.sub(lambda m: t.get(m.group(0), m.group(0)), str(text))
    left = sorted(set(TOKEN_RE.findall(out)))
    if left:
        die("%s 의 파일에서 치환어 %s 를 채울 값이 없다 — 배정 표 칸(재난안전과 계정·부서 계정·내 부서)과 "
            "기본값(정기평가 진행자·진행자 몫 부서)을 확인하라" % (person, ", ".join(left)))
    return out

def lead_only(r):
    return r[1].startswith(LEAD_TAG)

def who_of(r):
    return r[1][len(LEAD_TAG):].strip() if lead_only(r) else r[1]

def sheets_for(f, person):
    """그 사람이 받는 줄 — 진행자 줄은 진행자 파일에만"""
    return [(s, [r for r in s["rows"] if not lead_only(r) or person == LEAD]) for s in f["sheets"]]

def file_accounts(f):
    for t, blocks in f["guide"]:
        if ACCOUNTS_SECTION in t:
            for k, v in blocks:
                if k == "table" and v and v[0][0] == "이름":
                    return [r[0] for r in v[1:]]
    return ACCOUNTS

def gkey(t):
    return t.split()[0]

def guide_for(f):
    """공통 안내 + 파일별 안내. 첫 낱말이 같은 절은 바꾸고, 없는 절은 바로 앞 파일 절 뒤에 더한다."""
    fo = f["guide"]
    if not fo:
        return guide
    gkeys = {gkey(t) for t, _ in guide}
    at_key = {gkey(t): k for k, (t, _) in enumerate(fo)}
    out, used = [], set()
    def run_new(j):
        while j < len(fo) and gkey(fo[j][0]) not in gkeys:
            out.append(fo[j]); used.add(j); j += 1
    for gi, (t, b) in enumerate(guide):
        k = at_key.get(gkey(t))
        if k is None:
            out.append((t, b))
        else:
            out.append(fo[k]); used.add(k); run_new(k + 1)
        if gi == 0:
            run_new(0)        # 파일 절 맨 앞의 새 절은 공통 첫 절(하는 법) 바로 뒤에
    out += [fo[k] for k in range(len(fo)) if k not in used]
    return out

# 실제 메뉴 경로 — «어디서»가 화면과 다른 이름을 쓰면 테스터가 메뉴를 못 찾는다
def nav_paths():
    s = io.open(NAV_JS, encoding="utf-8").read()
    m = re.search(r"(?:const|var|let)\s+NAV\s*=\s*\[", s)
    if not m:
        die("js/layout.js 에서 NAV 를 찾지 못했다")
    k, depth = m.end() - 1, 0
    for j in range(k, len(s)):
        if s[j] == "[":
            depth += 1
        elif s[j] == "]":
            depth -= 1
            if depth == 0:
                break
    body, group, items = s[k:j], None, []
    for line in body.split("\n"):
        g = re.search(r"\{\s*id:\s*'[^']+',\s*label:\s*'([^']+)'.*items:\s*\[", line)
        if g:
            group = g.group(1); continue
        it = re.search(r"\{\s*id:\s*'[^']+'(.*?)label:\s*'([^']+)'", line)
        if it and group:
            sec_m = re.search(r"section:\s*'([^']+)'", line)
            items.append((group, sec_m.group(1) if sec_m else None, it.group(2), "hidden: true" in line))
    valid, short = set(), {}
    for grp, sct, lab, hidden in items:
        if hidden:
            continue
        short.setdefault("%s > %s" % (grp, lab), []).append(sct)
        if sct:
            valid.add("%s > %s > %s" % (grp, sct, lab))
    for p, secs in short.items():
        if len(secs) == 1:
            valid.add(p)       # 같은 이름이 둘이면(정기교육 등) 묶음 이름까지 적어야 한다
    if len(valid) < 20:
        die("js/layout.js NAV 에서 메뉴 경로를 %d개밖에 읽지 못했다 — 파서가 NAV 모양을 놓쳤다" % len(valid))
    return valid

VALID_PATHS = nav_paths()

# ── ② 원본 검사 ────────────────────────────────────────────────────────────
if not files:
    die("파일 절(## 파일: …)을 하나도 찾지 못했다")
names = [a[0] for a in assign]
if len(names) != len(set(names)) or not names:
    die("배정 표의 이름이 비었거나 중복됐다: %s" % names)
smoke = [f["name"] for f in files if f["kind"] == "스모크"]
given = [m for _, ms in assign for m in ms]
missing = [m for m in smoke if m not in given]
extra = [m for m in given if m not in smoke]
dup = sorted({m for m in given if given.count(m) > 1})
if missing or extra or dup:
    die("배정 표가 스모크 파일과 맞지 않는다 — 빠짐 %s · 없는 메뉴 %s · 두 번 %s" % (missing, extra, dup))
smoke_of = {n: ms for n, ms in assign}
def people_of(f):
    return [p for p in names if (f["kind"] == "상세" and not no_detail(p)) or f["name"] in smoke_of[p]]
if LEAD and LEAD not in names:
    die("기본값 «정기평가 진행자» «%s» 가 배정 표 이름이 아니다" % LEAD)
for f in files:
    nums = []
    accts = file_accounts(f)
    if not f["sheets"]:
        die("«%s» 에 시트가 없다" % f["name"])
    for s in f["sheets"]:
        if not s["rows"]:
            die("«%s / %s» 에 줄이 없다" % (f["name"], s["name"]))
        if len(s["name"]) > 31 or re.search(r"[\[\]:*?/\\]", s["name"]):
            die("시트 이름 «%s» 은 엑셀 규칙(31자·특수문자) 밖이다" % s["name"])
        if s["name"] in ("안내", "요약", "_정보"):
            die("시트 이름 «%s» 은 예약어다" % s["name"])
        for r in s["rows"]:
            tid, who, where = r[0], r[1], r[2]
            m = re.match(r"^%s-(\d{3})$" % f["code"], tid)
            if not m:
                die("«%s» 의 번호 «%s» 가 %s-000 형식이 아니다" % (f["name"], tid, f["code"]))
            n = int(m.group(1))
            if nums and n <= nums[-1]:
                die("«%s» 의 번호가 커지는 순서가 아니다: %s" % (f["name"], tid))
            nums.append(n)
            if not all(r[1:6]):
                die("%s 에 빈 칸이 있다(누구로·어디서·확인할 것·해 볼 것·이러면 성공은 모두 채운다)" % tid)
            for w in [x.strip() for x in who_of(r).split("→")]:
                if w not in accts:
                    die("%s «누구로»의 «%s» 가 계정 표(%s)에 없다" % (tid, w, ", ".join(accts)))
            if lead_only(r) and not LEAD:
                die("%s 는 진행자 줄인데 기본값에 «정기평가 진행자»가 없다" % tid)
            if " > " in where and where not in VALID_PATHS:
                die("%s «어디서» «%s» 가 실제 메뉴 경로가 아니다 — js/layout.js NAV 를 확인하라" % (tid, where))
    # 치환어 뒤 조사 — 원본 전체(줄·설명·파일 안내)
    texts = [(r[0], x) for s in f["sheets"] for r in s["rows"] for x in r[1:]]
    texts += [("설명", x) for x in f["desc"] + [d for s in f["sheets"] for d in s["desc"]]]
    texts += [("안내 " + t, v if k == "p" else " ".join(" ".join(row) for row in v))
              for t, blocks in f["guide"] for k, v in blocks]
    for where, x in texts:
        m = JOSA_RE.search(x)
        if m:
            die("«%s» %s — 치환어 «%s» 뒤에 조사 «%s» 가 붙어 있다. 값마다 받침이 달라 틀린 조사가 된다 "
                "— «박안전N 계정으로», «내 부서({내 부서})를» 처럼 쓴다" % (f["name"], where, m.group(1), m.group(2)))
    # 사람마다 — 진행자 줄을 뺀 뒤에도 시트가 비지 않고 «N번»이 그 사람 파일의 앞 줄을 가리키는가
    for p in people_of(f):
        got = sheets_for(f, p)
        pn = {int(r[0].split("-")[1]) for _, rows in got for r in rows}
        for s, rows in got:
            if not rows:
                die("«%s / %s» 가 %s 에게는 빈 시트다(진행자 줄만 있다)" % (f["name"], s["name"], p))
            for r in rows:
                cur = int(r[0].split("-")[1])
                for txt in r[4:6]:
                    for mm in re.finditer(r"(\d+)번(?!\s*시트)", txt):
                        ref = int(mm.group(1))
                        if ref not in pn or ref >= cur:
                            die("%s 가 가리키는 «%d번» 이 %s 파일의 앞 줄이 아니다(진행자 줄이면 다른 사람 파일에는 없다)" % (r[0], ref, p))
        # 치환 값이 다 있는가 — 쓰기 전에 멈춘다
        for s, rows in got:
            for r in rows:
                for x in r:
                    subst(x, p)
            for d in s["desc"]:
                subst(d, p)
        for t, blocks in guide_for(f):
            for k, v in blocks:
                subst(v if k == "p" else " ".join(" ".join(row) for row in v), p)

total = {f["name"]: sum(len(s["rows"]) for s in f["sheets"]) for f in files}
print("QA 엑셀 원본 — %s" % os.path.relpath(SRC, ROOT))
for f in files:
    lead_n = sum(1 for s in f["sheets"] for r in s["rows"] if lead_only(r))
    print("  %-10s %s %3d줄 · 시트 %d%s" % (f["name"], f["kind"], total[f["name"]], len(f["sheets"]),
          " (진행자 %s 만: %d줄 — 나머지는 %d줄)" % (LEAD, lead_n, total[f["name"]] - lead_n) if lead_n else ""))
print("  결과 값: %s" % " · ".join(RESULT_NAMES))
print("  인원 %d명: %s" % (len(names), ", ".join("%s(%s)" % (n, "·".join(ms)) for n, ms in assign)))
if CHECK:
    print("\n✔ 원본 형식 검사 통과 (--check — 파일을 쓰지 않았다)"); sys.exit(0)

# ── ③ 엑셀 쓰기 ────────────────────────────────────────────────────────────
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule, DataBarRule
from openpyxl.utils import get_column_letter

FONT = "맑은 고딕"          # 한글 문서의 표준 글꼴 — Arial 은 한글 글리프가 없어 섞여 보인다
f_base = Font(name=FONT, size=10.5)
f_bold = Font(name=FONT, size=10.5, bold=True)
f_head = Font(name=FONT, size=10.5, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=15, bold=True)
f_banner = Font(name=FONT, size=10.5, bold=True, color="1F4E79")
f_sec = Font(name=FONT, size=12, bold=True, color="1F4E79")
f_gray = Font(name=FONT, size=10.5, color="7F7F7F", italic=True)
f_dim = Font(name=FONT, size=10.5, color="A6A6A6")
fill_head = PatternFill("solid", fgColor="1F4E79")
fill_head_in = PatternFill("solid", fgColor="C55A11")      # 테스터가 채우는 칸의 머리 — 색으로 «여기만»
fill_input = PatternFill("solid", fgColor="FFF2CC")
fill_key = PatternFill("solid", fgColor="F2F2F2")
fill_banner = PatternFill("solid", fgColor="DDEBF7")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

HEAD = ["번호", "누구로", "어디서", "확인할 것", "해 볼 것", "이러면 성공",
        "결과 ▼", "메모 (실패·오류·문의일 때)", "사진 (필요할 때 붙여넣기)"]
WIDTH = [5, 14, 24, 20, 38, 38, 10, 26, 22]
COL_RESULT, COL_MEMO, COL_PHOTO = 7, 8, 9
HEAD_ROW, FIRST_ROW = 3, 4

def char_w(ch):
    return 1.0 if ord(ch) < 0x1100 else 1.85

def est_lines(text, width):
    tot = 0
    for part in str(text).split("\n"):
        w = sum(char_w(c) for c in part)
        tot += max(1, -(-int(w) // max(1, int(width * 1.1))))
    return tot

def fit_row(ws, r, texts_widths, min_h=20):
    n = max(est_lines(t, w) for t, w in texts_widths) if texts_widths else 1
    ws.row_dimensions[r].height = max(min_h, min(409, 15 * n + 6))

def steps_text(s):
    # «① A ② B» → 줄마다 한 단계. 원본은 한 줄로 쓰고 엑셀에서 읽기 쉽게 편다.
    s = re.sub(r"\s*([②③④⑤⑥⑦⑧⑨⑩])", r"\n\1", s.strip())
    return s.replace("<br>", "\n")

def clean(s):
    return s.replace("**", "").replace("`", "")

def q(sheet):
    return "'" + sheet.replace("'", "''") + "'"

def result_rules(ws, rng):
    for name, _, color in RESULTS:
        bg, fg = COLOR[color]
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"%s"' % name],
            fill=PatternFill("solid", fgColor=bg), font=Font(name=FONT, size=10.5, color=fg, bold=True)))

def result_dv(ranges):
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(RESULT_NAMES), allow_blank=True,
                        showInputMessage=True, promptTitle="결과 고르기",
                        prompt=" · ".join(RESULT_NAMES) + " 중에서 고르세요. 실패·오류·문의면 옆 메모에 한 줄.",
                        showErrorMessage=True, errorTitle="목록에서 고르세요",
                        error="결과는 " + " · ".join(RESULT_NAMES) + " 중 하나입니다.")
    for rg in ranges:
        dv.add(rg)
    return dv

def build(f, person, gen_date):
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True      # 열 때 요약 수식을 다시 계산한다
    menu, code = f["name"], f["code"]
    fname = "%s %s %s.xlsx" % (PREFIX, menu, person)
    sheet_names = [s["name"] for s in f["sheets"]]
    got = sheets_for(f, person)
    n_rows = sum(len(rows) for _, rows in got)
    uses_acct = any(TOKEN_RE.search(x) for s in f["sheets"] for r in s["rows"] for x in r[1:2])
    lead_n = sum(1 for s in f["sheets"] for r in s["rows"] if lead_only(r))

    # ── 안내 ──
    ws = wb.active; ws.title = "안내"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "595959"
    ws.column_dimensions["A"].width = 18
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 17
    ws.cell(row=1, column=1, value="%s — %s · %s" % (PREFIX, menu, person)).font = f_title
    ws.merge_cells("A1:H1"); ws.row_dimensions[1].height = 30
    r = 3
    url = meta.get("테스트 주소", "")
    info = [("메뉴", menu + (" (자세히 보는 파일)" if f["kind"] == "상세" else " (훑어보는 파일)"), False),
            ("담당자", person, False)]
    if uses_acct:
        info.append(("내 계정", subst("박안전N (재난안전과 담당자) · 하정수N (내 부서 {내 부서} 담당 공무원)", person), False))
    if lead_n:
        info.append(("정기평가 진행자", "나 — 진행자만 하는 줄 %d개가 이 파일에 들어 있습니다(«누구로»에 «(진행자)»)" % lead_n
                     if person == LEAD else "%s — 진행자만 하는 줄 %d개는 진행자 파일에만 있어서 번호가 건너뜁니다" % (LEAD, lead_n), False))
    info += [("이 파일", subst(" ".join(f["desc"]), person), False),
            ("줄 수", "%d줄" % n_rows, False),
            ("시트", " → ".join(sheet_names), False),
            ("테스트 주소", url, True),
            ("브라우저", None, True),
            ("기기", None, True),
            ("시작한 날", None, True)]
    dv_browser = DataValidation(type="list", formula1='"%s"' % ",".join(BROWSERS), allow_blank=True)
    dv_device = DataValidation(type="list", formula1='"%s"' % ",".join(DEVICES), allow_blank=True)
    for k, v, is_input in info:
        a = ws.cell(row=r, column=1, value=k); a.font = f_bold; a.fill = fill_key; a.border = border; a.alignment = wrap
        b = ws.cell(row=r, column=2, value=v); b.font = f_base; b.border = border; b.alignment = wrap
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        if is_input:
            b.fill = fill_input
        if k == "테스트 주소":
            if str(v).startswith("http"):
                b.hyperlink = v; b.font = Font(name=FONT, size=10.5, color="0563C1", underline="single")
            else:
                b.font = f_gray
        if k == "브라우저":
            dv_browser.add(b.coordinate)
        if k == "기기":
            dv_device.add(b.coordinate)
        if k == "시작한 날":
            b.number_format = "yyyy-mm-dd"
        fit_row(ws, r, [(v or "", 17 * 7)], 22)
        r += 1
    ws.add_data_validation(dv_browser); ws.add_data_validation(dv_device)
    r += 1
    for sec_title, blocks in guide_for(f):
        sec_title = subst(sec_title, person)
        blocks = [(k, subst(v, person) if k == "p" else [[subst(x, person) for x in row] for row in v]) for k, v in blocks]
        ws.cell(row=r, column=1, value=sec_title).font = f_sec
        ws.row_dimensions[r].height = 24; r += 1
        for kind, val in blocks:
            if kind == "p":
                txt = re.sub(r"^[-*]\s+", "", val)
                txt = clean(txt if re.match(r"^[①-⑩]", txt) else "• " + txt)
                c = ws.cell(row=r, column=1, value=txt); c.font = f_base; c.alignment = wrap
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
                fit_row(ws, r, [(txt, 18 + 17 * 7)]); r += 1
                continue
            is_result = sec_title == RESULTS_SECTION
            rows = [x[:2] for x in val] if is_result else val     # 결과 표의 «색» 칸은 색으로만 보여 준다
            ncol = max(len(x) for x in rows)
            spans = [(1, 1)]
            if ncol > 1:
                rest = ncol - 1; each = max(1, 7 // rest); c0 = 2
                for ci in range(rest):
                    c1 = 8 if ci == rest - 1 else min(8, c0 + each - 1)
                    spans.append((c0, c1)); c0 = c1 + 1
            for ri, row in enumerate(rows):
                for ci, v in enumerate(row):
                    s0, s1 = spans[ci]
                    c = ws.cell(row=r, column=s0, value=clean(v))
                    c.border = border; c.alignment = wrap
                    c.font = f_bold if (ri == 0 or ci == 0) else f_base
                    if ri == 0:
                        c.fill = fill_key
                    elif is_result and ci == 0:
                        bg, fg = COLOR[val[ri][2]]
                        c.fill = PatternFill("solid", fgColor=bg)
                        c.font = Font(name=FONT, size=10.5, bold=True, color=fg)
                        c.alignment = center
                    if s1 > s0:
                        ws.merge_cells(start_row=r, start_column=s0, end_row=r, end_column=s1)
                fit_row(ws, r, [(v, 18 if ci == 0 else 17 * (spans[ci][1] - spans[ci][0] + 1))
                                for ci, v in enumerate(row)], 22)
                r += 1
        r += 1

    # ── 시나리오 시트 ──
    ranges = []
    for s, srows in got:
        wt = wb.create_sheet(s["name"])
        wt.sheet_properties.tabColor = "70AD47"
        t = wt.cell(row=1, column=1, value="%s — %s" % (menu, s["name"])); t.font = f_title
        wt.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEAD))
        wt.row_dimensions[1].height = 28
        banner = subst(" ".join(s["desc"]), person) or "위에서부터 차례로 하세요."
        b = wt.cell(row=2, column=1, value="🗂 " + banner); b.font = f_banner; b.fill = fill_banner; b.alignment = wrap
        wt.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEAD))
        fit_row(wt, 2, [(banner, sum(WIDTH))], 24)
        for k, h in enumerate(HEAD, 1):
            c = wt.cell(row=HEAD_ROW, column=k, value=h)
            c.font = f_head; c.alignment = center; c.border = border
            c.fill = fill_head_in if k >= COL_RESULT else fill_head
            wt.column_dimensions[get_column_letter(k)].width = WIDTH[k - 1]
        wt.row_dimensions[HEAD_ROW].height = 30
        r0 = FIRST_ROW
        shown = []
        for x in srows:
            who = " → ".join(subst(w, person) + ("(%s)" % ROLE_OF[w] if w in ROLE_OF else "")
                             for w in [y.strip() for y in who_of(x).split("→")])
            who += ("\n진행자만" if lead_only(x) else "")
            shown.append([x[0], who] + [subst(v, person) for v in x[2:]])
        for k, x in enumerate(shown):
            rr = r0 + k
            vals = [int(x[0].split("-")[1]), x[1], x[2], clean(x[3]),
                    clean(steps_text(x[4])), clean(x[5]), None, None, None]
            for col, v in enumerate(vals, 1):
                c = wt.cell(row=rr, column=col, value=v)
                c.border = border; c.font = f_base
                c.alignment = center if col in (1, COL_RESULT) else wrap
                if col >= COL_RESULT:
                    c.fill = fill_input
            wt.cell(row=rr, column=1).font = f_bold
            wt.cell(row=rr, column=4).font = f_bold
            if k and shown[k - 1][1] == x[1]:
                wt.cell(row=rr, column=2).font = f_dim
            if k and shown[k - 1][2] == x[2]:
                wt.cell(row=rr, column=3).font = f_dim
            wt.cell(row=rr, column=COL_RESULT).font = Font(name=FONT, size=10.5, bold=True)
            fit_row(wt, rr, [(vals[j], WIDTH[j]) for j in (1, 2, 3, 4, 5)], 34)
        r1 = r0 + len(shown) - 1
        H, I = get_column_letter(COL_RESULT), get_column_letter(COL_MEMO)
        wt.add_data_validation(result_dv(["%s%d:%s%d" % (H, r0, H, r1)]))
        result_rules(wt, "%s%d:%s%d" % (H, r0, H, r1))
        # 실패·오류·문의인데 메모가 비면 메모 칸을 붉게 — «왜»가 없는 실패는 고칠 수 없다
        cond = "OR(%s)" % ",".join('$%s%d="%s"' % (H, r0, x) for x in FAIL_LIKE)
        wt.conditional_formatting.add("%s%d:%s%d" % (I, r0, I, r1), FormulaRule(
            formula=['AND(%s,$%s%d="")' % (cond, I, r0)], fill=PatternFill("solid", fgColor="FF9E9E")))
        # 번호는 옆으로 밀어도 보이게 — 결과 칸에서 몇 번 줄인지 잃지 않는다
        wt.freeze_panes = "B%d" % FIRST_ROW
        wt.sheet_view.zoomScale = 90
        wt.auto_filter.ref = "A%d:%s%d" % (HEAD_ROW, get_column_letter(len(HEAD)), r1)
        wt.page_setup.orientation = "landscape"; wt.page_setup.fitToWidth = 1; wt.page_setup.fitToHeight = 0
        wt.sheet_properties.pageSetUpPr.fitToPage = True
        wt.print_title_rows = "%d:%d" % (HEAD_ROW, HEAD_ROW)
        ranges.append((s["name"], r0, r1))

    # ── 요약 (수식) ──
    wsum = wb.create_sheet("요약")
    wsum.sheet_view.showGridLines = False
    wsum.sheet_properties.tabColor = "2F75B5"
    heads = ["시트", "전체"] + RESULT_NAMES + ["남은 것", "진행률"]
    wsum.cell(row=1, column=1, value="진행 상황 — %s · %s" % (menu, person)).font = f_title
    wsum.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(heads))
    wsum.row_dimensions[1].height = 30
    note = wsum.cell(row=2, column=1, value="결과를 고르면 저절로 채워집니다. 남은 것 = 아직 결과를 고르지 않은 줄입니다.")
    note.font = f_gray
    wsum.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(heads))
    widths = [24, 8] + [9] * len(RESULT_NAMES) + [9, 12]
    for k, h in enumerate(heads, 1):
        c = wsum.cell(row=4, column=k, value=h); c.font = f_head; c.fill = fill_head; c.alignment = center; c.border = border
        wsum.column_dimensions[get_column_letter(k)].width = widths[k - 1]
    wsum.row_dimensions[4].height = 26
    cached, rr = {}, 5
    nres = len(RESULT_NAMES)
    c_left, c_prog = 3 + nres, 4 + nres
    L = get_column_letter
    for name, r0, r1 in ranges:
        Hr = "%s!$%s$%d:$%s$%d" % (q(name), L(COL_RESULT), r0, L(COL_RESULT), r1)
        Ar = "%s!$A$%d:$A$%d" % (q(name), r0, r1)
        row = [name, "=COUNTA(%s)" % Ar] + ['=COUNTIF(%s,"%s")' % (Hr, v) for v in RESULT_NAMES]
        row += ["=COUNTBLANK(%s)" % Hr,
                "=IF(B{0}=0,0,(B{0}-{1}{0})/B{0})".format(rr, L(c_left))]
        init = [r1 - r0 + 1] + [0] * nres + [r1 - r0 + 1, 0]
        for col, v in enumerate(row, 1):
            c = wsum.cell(row=rr, column=col, value=v); c.border = border; c.font = f_base
            c.alignment = Alignment(horizontal="left" if col == 1 else "center", vertical="center")
            if col >= 2:
                cached["%s%d" % (L(col), rr)] = init[col - 2]
        wsum.cell(row=rr, column=c_prog).number_format = "0%"
        wsum.row_dimensions[rr].height = 22
        rr += 1
    first, last = 5, rr - 1
    tot = ["합계"] + ["=SUM(%s%d:%s%d)" % (L(c), first, L(c), last) for c in range(2, c_prog)]
    tot += ["=IF(B{0}=0,0,(B{0}-{1}{0})/B{0})".format(rr, L(c_left))]
    for col, v in enumerate(tot, 1):
        c = wsum.cell(row=rr, column=col, value=v); c.border = border; c.font = f_bold; c.fill = fill_key
        c.alignment = Alignment(horizontal="left" if col == 1 else "center", vertical="center")
        if col >= 2:
            cached["%s%d" % (L(col), rr)] = 0 if col == c_prog else sum(cached["%s%d" % (L(col), k)] for k in range(first, last + 1))
    wsum.cell(row=rr, column=c_prog).number_format = "0%"
    wsum.row_dimensions[rr].height = 24
    prog = "%s%d:%s%d" % (L(c_prog), first, L(c_prog), rr)
    wsum.conditional_formatting.add(prog, DataBarRule(start_type="num", start_value=0, end_type="num",
                                                      end_value=1, color="63BE7B", showValue=True))
    # 실패·오류가 한 건이라도 있으면 숫자를 붉게 — 요약만 봐도 어디를 봐야 할지 보이게
    for x in FAIL_LIKE:
        col = 3 + RESULT_NAMES.index(x)
        wsum.conditional_formatting.add("%s%d:%s%d" % (L(col), first, L(col), rr), CellIsRule(
            operator="greaterThan", formula=["0"], font=Font(name=FONT, size=10.5, bold=True, color="C00000")))

    # ── _정보 (숨김) — 취합 스크립트가 읽는 값. 파일 이름이 바뀌어도 누구의 무엇인지 안다 ──
    wi = wb.create_sheet("_정보")
    kv = [("형식", FORMAT), ("메뉴", menu), ("담당자", person), ("파일 코드", code), ("종류", f["kind"]),
          ("생성일", gen_date), ("결과 값", ",".join(RESULT_NAMES)), ("머리 행", HEAD_ROW),
          ("열", ",".join(h.split(" ")[0] for h in HEAD))]
    for k, (a, b) in enumerate(kv, 1):
        wi.cell(row=k, column=1, value=a); wi.cell(row=k, column=2, value=b)
    wi.sheet_state = "hidden"

    # 탭 순서 — 안내 · 요약 · 시나리오 … · (숨김) _정보
    wb._sheets = [wb[n] for n in ["안내", "요약"] + sheet_names + ["_정보"]]
    wb.active = 0
    return fname, wb, cached

def inject_cached(path, sheet_title, values):
    """openpyxl 은 수식의 결과값(캐시)을 비워 둔다. 엑셀은 열 때 다시 계산하지만(fullCalcOnLoad),
    메일·메신저로 받은 파일을 여는 «제한된 보기»와 미리보기(Finder·드라이브)는 캐시만 보여 줘
    요약이 전부 0·빈칸으로 보인다. 처음 상태(전원 미입력)의 값을 캐시로 넣어 둔다 —
    결과를 입력하면 엑셀이 다시 계산하므로 값이 굳지 않는다."""
    import zipfile, html
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    by = {i.filename: d for i, d in items}
    wbx = by["xl/workbook.xml"].decode("utf-8")
    rels = by["xl/_rels/workbook.xml.rels"].decode("utf-8")
    rid = None
    for m in re.finditer(r'<sheet\b[^>]*>', wbx):
        tag = m.group(0)
        if html.unescape(re.search(r'name="([^"]*)"', tag).group(1)) == sheet_title:
            rid = re.search(r'r:id="([^"]*)"', tag).group(1)
    if not rid:
        die("캐시 주입 — «%s» 시트를 workbook.xml 에서 찾지 못했다" % sheet_title)
    target = None
    for m in re.finditer(r'<Relationship\b[^>]*>', rels):
        tag = m.group(0)
        if re.search(r'Id="%s"' % re.escape(rid), tag):
            target = re.search(r'Target="([^"]*)"', tag).group(1).lstrip("/")
    if not target:
        die("캐시 주입 — %s 의 대상 파일을 찾지 못했다" % rid)
    if not target.startswith("xl/"):
        target = "xl/" + target
    out, hit = [], 0
    for info, data in items:
        if info.filename == target:
            xml = data.decode("utf-8")
            def fill(m):
                nonlocal hit
                if m.group(1) not in values:
                    return m.group(0)
                hit += 1
                return m.group(0).replace("<v></v>", "<v>%s</v>" % values[m.group(1)])
            xml = re.sub(r'<c r="([A-Z]+\d+)"[^>]*><f>.*?</f><v></v></c>', fill, xml)
            data = xml.encode("utf-8")
        out.append((info, data))
    if hit != len(values):
        # 조용히 빗나가면 요약이 일부만 채워진 채 배포된다 — 멈춘다
        die("캐시 주입 — %d개 중 %d개만 찾았다(%s)" % (len(values), hit, os.path.basename(path)))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in out:
            z.writestr(info, data)

if __name__ == "__main__":
    gen_date = datetime.date.today().isoformat()
    made = []
    os.makedirs(OUT, exist_ok=True)
    plan = [(f, p) for p in names for f in files if p in people_of(f)]
    for f, person in plan:
        fname, wb, cached = build(f, person, gen_date)
        d = os.path.join(OUT, person)
        os.makedirs(d, exist_ok=True)
        path = os.path.join(d, fname)
        wb.save(path)
        inject_cached(path, "요약", cached)
        made.append(os.path.relpath(path, ROOT))
    # 지난 생성에서 남은 파일(배정이 바뀐 경우)을 알린다 — 지우지는 않는다.
    # macOS 는 한글 파일명을 NFD 로 돌려줄 때가 있어 정규화해 비교한다.
    made_n = {nfc(p) for p in made}
    stale = []
    for dp, _, fs in os.walk(OUT):
        for fn in fs:
            rel = os.path.relpath(os.path.join(dp, fn), ROOT)
            if fn.endswith(".xlsx") and not fn.startswith("~$") and nfc(rel) not in made_n:
                stale.append(rel)
    print("\n✔ %d개 파일 생성 → %s" % (len(made), os.path.relpath(OUT, ROOT)))
    for p in made:
        print("  · " + p)
    if stale:
        print("\n⚠ 이번 원본에 없는 옛 파일 %d개 — 배정이 바뀌었으면 직접 지운다:" % len(stale))
        for p in stale:
            print("  · " + p)
